import csv
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

DB_TABLE = "MatchData"
OUTPUT_DOCX = "output.docx"
OUTPUT_XLSX = "output.xlsx"

def load_from_db():
    import psycopg2
    conn = psycopg2.connect(
        database="Grafana-Output",
        host="127.0.0.1",
        user="postgres",
        password="MA6328",
        port="23010",
    )
    cur = conn.cursor()
    cur.execute(f'SELECT * FROM "{DB_TABLE}"')
    data = cur.fetchall()
    col_names = [desc[0] for desc in cur.description]
    cur.close()
    conn.close()
    return col_names, [list(row) for row in data]


def load_from_csv(csv_path):
    with open(csv_path, newline="", encoding="utf-8") as f:
        reader = csv.reader(f)
        rows = list(reader)
    return rows[0], rows[1:]



def export_xlsx(col_names, data, output_path):
    wb = Workbook()
    ws = wb.active
    ws.title = DB_TABLE

    header_fill = PatternFill("solid", start_color="2E75B6")
    header_font = Font(bold=True, color="FFFFFF", name="Arial", size=11)
    cell_font   = Font(name="Arial", size=10)

    ws.append(col_names)
    for col_idx, _ in enumerate(col_names, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font      = header_font
        cell.fill      = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in data:
        ws.append([str(v) for v in row])

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = cell_font

    for col_cells in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col_cells), default=10)
        ws.column_dimensions[col_cells[0].column_letter].width = min(max_len + 4, 40)

    ws.freeze_panes = "A2"
    wb.save(output_path)
    print(f"XLSX saved - {output_path}")

def set_cell_bg(cell, hex_color):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement("w:shd")
    shd.set(qn("w:val"),   "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"),  hex_color)
    tcPr.append(shd)


def export_docx(col_names, data, output_path):
    doc = Document()

    for section in doc.sections:
        section.top_margin    = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin   = Inches(1)
        section.right_margin  = Inches(1)

    title = doc.add_heading("Match Data", level=1)
    title.runs[0].font.color.rgb = RGBColor(0, 0, 0)
    title.runs[0].font.name      = "Arial"

    table = doc.add_table(rows=1, cols=len(col_names))
    table.style = "Table Grid"

    hdr_cells = table.rows[0].cells
    for i, col in enumerate(col_names):
        cell = hdr_cells[i]
        cell.text = col
        run = cell.paragraphs[0].runs[0]
        run.bold           = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        run.font.name      = "Arial"
        run.font.size      = Pt(10)
        cell.paragraphs[0].alignment = 1
        set_cell_bg(cell, "2E75B6")

    for row_data in data:
        row_cells = table.add_row().cells
        for i, val in enumerate(row_data):
            cell = row_cells[i]
            cell.text = str(val)
            run = cell.paragraphs[0].runs[0]
            run.font.name = "Arial"
            run.font.size = Pt(9)

    doc.save(output_path)
    print(f"DOCX saved - {output_path}")

def export_all(source="db", csv_path=None):
    if source == "db":
        print("Loading from database…")
        col_names, data = load_from_db()
    elif source == "csv":
        print(f"Loading from CSV: {csv_path}")
        col_names, data = load_from_csv(csv_path)
    else:
        raise ValueError("source must be 'db' or 'csv'")

    print(f"Loaded {len(data)} rows x {len(col_names)} columns")
    export_xlsx(col_names, data, OUTPUT_XLSX)
    export_docx(col_names, data, OUTPUT_DOCX)


if __name__ == "__main__":
    #export_all(source="db")
    export_all(source="csv", csv_path="file.csv")